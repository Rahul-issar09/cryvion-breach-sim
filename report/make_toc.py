"""Standalone Table of Contents document to paste into the main report."""
import os
from docx import Document
from docx.shared import Pt, Inches, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_LINE_SPACING
from docx.oxml.ns import qn

FONT = "Times New Roman"
DARK = RGBColor(0x2c, 0x3e, 0x50)
RIGHT_TAB = Inches(5.9)
OUT = os.path.join(os.path.dirname(__file__), "Table_of_Contents.docx")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = Inches(1.3)
sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1.0)

n = doc.styles["Normal"]
n.font.name = FONT
n.font.size = Pt(12)
n.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
n.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE


def run(p, text, bold=False, size=12):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    return r


def entry(title, page="", bold=False, indent=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run(p, title, bold=bold)
    run(p, "\t" + str(page), bold=bold)


# --- heading ---
h = doc.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
h.paragraph_format.space_after = Pt(12)
rr = run(h, "TABLE OF CONTENTS", bold=True, size=14)
rr.font.color.rgb = DARK

# --- preliminary pages (roman) ---
for t, pg in [("Abstract", "i"), ("List of Figures", "ii"),
              ("List of Tables", "iii"), ("List of Abbreviations", "iv")]:
    entry(t, pg, bold=True)
doc.add_paragraph().paragraph_format.space_after = Pt(2)

# --- chapters ---
chapters = [
    ("1", "Introduction", [
        ("1.1", "Background"), ("1.2", "Motivation"), ("1.3", "Objectives"),
        ("1.4", "Scope and Limitations"), ("1.5", "Organization of the Report")]),
    ("2", "Literature Review", [
        ("2.1", "The Cyber Kill Chain"), ("2.2", "The MITRE ATT&CK Framework"),
        ("2.3", "NIST SP 800-61 Incident-Response Lifecycle"),
        ("2.4", "Container Security Fundamentals"),
        ("2.5", "SIEM and Host-Based Detection"),
        ("2.6", "Web Application Security and the WAF"),
        ("2.7", "Zero Trust and Micro-Segmentation"),
        ("2.8", "DevSecOps Tooling"), ("2.9", "Related Work and Research Gap")]),
    ("3", "Problem Identification", [
        ("3.1", "The Core Problem"), ("3.2", "Specific Gaps"),
        ("3.3", "Problem Statement"), ("3.4", "Derived Requirements")]),
    ("4", "Methodology", [
        ("4.1", "System Architecture"), ("4.2", "Vulnerable Environment Design"),
        ("4.3", "Attack Simulation"), ("4.4", "Detection and Monitoring"),
        ("4.5", "Forensic Methodology"), ("4.6", "Incident-Response Lifecycle"),
        ("4.7", "Hardened Re-Architecture"), ("4.8", "DevSecOps CI/CD Pipeline")]),
    ("5", "Results", [
        ("5.1", "Breach Execution"), ("5.2", "Detection Results"),
        ("5.3", "Forensic Findings"), ("5.4", "Hardening Validation"),
        ("5.5", "DevSecOps Results"), ("5.6", "Discussion")]),
    ("6", "Conclusion", [
        ("6.1", "Summary"), ("6.2", "Limitations"), ("6.3", "Future Work")]),
]
for num, title, subs in chapters:
    entry(f"Chapter {num}   {title}", "", bold=True)
    for s_num, s_title in subs:
        entry(f"{s_num}   {s_title}", "", indent=0.4)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# --- back matter ---
for t in ["Bibliography", "Appendix"]:
    entry(t, "", bold=True)

doc.save(OUT)
print("WROTE", OUT)
