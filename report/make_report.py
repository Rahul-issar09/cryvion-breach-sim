"""
Generate the CSVTU / RCET B.Tech (6th sem) Major Project Report as a .docx.
Title: Design and Evaluation of an End-to-End Breach Lifecycle Simulation in
Containerized Environments.  Fill the <PLACEHOLDER> fields, then insert screenshots.
"""
import os
from docx import Document
from docx.shared import Pt, Inches, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(__file__)
FIG = os.path.join(HERE, "figures")
OUT = os.path.join(HERE, "Cryvion_Major_Project_Report.docx")

FONT = "Times New Roman"
GREY = RGBColor(0x66, 0x66, 0x66)
DARK = RGBColor(0x2c, 0x3e, 0x50)

doc = Document()

# ---- page setup: A4, margins L1.3 R/T/B 1 ----
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = Inches(1.3)
sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1.0)

# ---- base styles ----
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
normal.paragraph_format.space_after = Pt(6)


def _set_font(run, size=12, bold=False, italic=False, smallcaps=False, color=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.small_caps = smallcaps
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), name)


# Configure Word's built-in Heading styles so the TOC field can read them,
# while matching the required look (TNR 14 UPPER for H1, TNR 14 small-caps for H2).
def _cfg_heading(name, size, smallcaps=False, italic=False):
    st = doc.styles[name]
    st.font.name = FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.italic = italic
    st.font.small_caps = smallcaps
    st.font.color.rgb = DARK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.keep_with_next = True


_cfg_heading("Heading 1", 14)
_cfg_heading("Heading 2", 14, smallcaps=True)
_cfg_heading("Heading 3", 12, italic=True)


def h1(text):
    """Chapter heading: TNR 14, ALL UPPERCASE, bold (Word Heading 1 -> TOC)."""
    return doc.add_paragraph(text.upper(), style="Heading 1")


def h2(text):
    """Sub-heading: TNR 14, small caps (Word Heading 2 -> TOC)."""
    return doc.add_paragraph(text, style="Heading 2")


def h3(text):
    return doc.add_paragraph(text, style="Heading 3")


def body(text, justify=True):
    p = doc.add_paragraph(text)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def numbered(text):
    return doc.add_paragraph(text, style="List Number")


def centered(text, size=12, bold=False, italic=False, smallcaps=False, color=None, space=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space)
    _set_font(p.add_run(text), size=size, bold=bold, italic=italic, smallcaps=smallcaps, color=color)
    return p


def figure(img, caption):
    if os.path.exists(img):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img, width=Inches(5.9))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(cap.add_run(caption), size=11, italic=True, color=GREY)
    cap.paragraph_format.space_after = Pt(10)


def screenshot(caption, hint):
    """A placeholder box for a screenshot the student inserts, with caption."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = t.cell(0, 0)
    cell.width = Inches(5.9)
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(cp.add_run("[ INSERT SCREENSHOT HERE ]\n" + hint), size=11, italic=True, color=GREY)
    cp.paragraph_format.space_before = Pt(24)
    cp.paragraph_format.space_after = Pt(24)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(cap.add_run(caption), size=11, italic=True, color=GREY)
    cap.paragraph_format.space_after = Pt(10)


def table(caption, headers, rows, widths=None):
    cap = doc.add_paragraph()
    _set_font(cap.add_run(caption), size=11, bold=True, color=DARK)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, htext in enumerate(headers):
        c = t.rows[0].cells[i]
        c.paragraphs[0].clear()
        _set_font(c.paragraphs[0].add_run(htext), size=10, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].clear()
            _set_font(cells[i].paragraphs[0].add_run(str(val)), size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    _set_font(p.add_run(text), size=9, name="Consolas")
    return p


def page_break():
    doc.add_page_break()


def _pgnum(section, fmt, start=None):
    sectPr = section._sectPr
    pgnt = sectPr.find(qn("w:pgNumType"))
    if pgnt is None:
        pgnt = OxmlElement("w:pgNumType")
        sectPr.append(pgnt)
    pgnt.set(qn("w:fmt"), fmt)
    if start is not None:
        pgnt.set(qn("w:start"), str(start))


def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
    _set_font(run, size=11)


def insert_toc():
    p = doc.add_paragraph()
    run = p.add_run()
    fb = OxmlElement("w:fldChar"); fb.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = r'TOC \o "1-3" \h \z \u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "Right-click → Update Field to build the Table of Contents."
    fe = OxmlElement("w:fldChar"); fe.set(qn("w:fldCharType"), "end")
    for e in (fb, it, sep, txt, fe):
        run._r.append(e)


# ===================================================================
# FRONT MATTER
# ===================================================================
def front_matter():
    # ---- Title page ----
    centered("A Major Project Report", 14, bold=True, space=10)
    centered("On", 12, space=10)
    centered("Design and Evaluation of an End-to-End Breach Lifecycle Simulation in Containerized Environments",
             14, bold=True, space=16)
    centered("Submitted to", 12, space=4)
    centered("CHHATTISGARH SWAMI VIVEKANAND TECHNICAL UNIVERSITY, BHILAI", 12, bold=True, space=14)
    centered("In partial fulfilment of the requirements for the award of the degree of", 12, space=4)
    centered("Bachelor of Technology", 13, bold=True, space=4)
    centered("In", 12, space=4)
    centered("<NAME OF THE BRANCH>", 12, bold=True, space=4)
    centered("(6th Semester)", 12, space=14)
    centered("By", 12, space=8)
    for i in range(1, 7):
        centered(f"<Student Name {i}, CSVTU Roll No. {i}, Enrollment No. {i}>", 11, space=2)
    centered("", space=10)
    centered("Under the Guidance of", 12, space=4)
    centered("<Name of Guide>", 12, bold=True, space=2)
    centered("<Designation>", 12, space=12)
    centered("Department of <Name of the Department>", 12, bold=True, space=2)
    centered("Rungta College of Engineering & Technology, Bhilai (C.G.)", 12, bold=True, space=2)
    centered("Session <20XX–20XX>", 12, space=2)
    page_break()

    # ---- Declaration ----
    h1("Declaration")
    body("We, the undersigned, solemnly declare that this report on the project work entitled "
         "“Design and Evaluation of an End-to-End Breach Lifecycle Simulation in Containerized "
         "Environments”, is based on our own work carried out during the course of our study under "
         "the guidance of <Guide’s Name>. We assert that the statements made and conclusions drawn are "
         "an outcome of the project work. We further declare that to the best of our knowledge and belief, "
         "the report does not contain any part of any work which has been submitted for the award of any "
         "other degree/diploma/certificate in this University or any other University.")
    doc.add_paragraph()
    for i in range(1, 7):
        centered(f"<Student Name {i}>   <CSVTU Roll No. {i}>   <Enrollment No. {i}>", 11, space=4)
    page_break()

    # ---- Certificate (Guide/HoD) ----
    h1("Certificate")
    body("This is to certify that this report on the project submitted is an outcome of the project work "
         "entitled “Design and Evaluation of an End-to-End Breach Lifecycle Simulation in Containerized "
         "Environments”, carried out by the students named in the Declaration, under my guidance and "
         "supervision for the award of the Degree of Bachelor of Technology in <Name of the Department> of "
         "Chhattisgarh Swami Vivekanand Technical University, Bhilai (C.G.), India.")
    body("To the best of my knowledge, the report embodies the work of the students themselves, has duly "
         "been completed, fulfils the requirement of the Ordinance relating to the B.Tech. degree of the "
         "University, and is up to the desired standard for the purpose for which it is submitted.")
    doc.add_paragraph(); doc.add_paragraph()
    centered("<Signature>", 12, space=2)
    centered("<Name of Guide>", 12, bold=True, space=2)
    centered("<Designation>", 12, space=12)
    body("This project work, as mentioned above, is hereby being recommended and forwarded for examination "
         "and evaluation by the University.")
    doc.add_paragraph()
    centered("<Signature>", 12, space=2)
    centered("<Name of HoD>", 12, bold=True, space=2)
    centered("Head, Department of <Name of the Department>", 12, space=2)
    centered("Rungta College of Engineering & Technology, Bhilai (C.G.), India", 12, space=2)
    page_break()

    # ---- Certificate by Examiners ----
    h1("Certificate by the Examiners")
    body("This is to certify that this project work entitled “Design and Evaluation of an End-to-End "
         "Breach Lifecycle Simulation in Containerized Environments”, submitted by the students named in "
         "the Declaration, is duly examined by the undersigned as a part of the examination for the award of "
         "the Bachelor of Technology degree in <Name of the Department> of Chhattisgarh Swami Vivekanand "
         "Technical University, Bhilai.")
    doc.add_paragraph(); doc.add_paragraph()
    p = doc.add_paragraph()
    _set_font(p.add_run("Internal Examiner\t\t\t\t\tExternal Examiner"), size=12)
    p2 = doc.add_paragraph()
    _set_font(p2.add_run("Name & Signature\t\t\t\t\tName & Signature"), size=12)
    p3 = doc.add_paragraph()
    _set_font(p3.add_run("Date:\t\t\t\t\t\tDate:"), size=12)
    page_break()

    # ---- Acknowledgements ----
    h1("Acknowledgements")
    body("It is a matter of profound privilege and pleasure to extend our sense of respect and deepest "
         "gratitude to our project guide <Guide’s Name>, Department of <Name of the Department>, under "
         "whose precise guidance and gracious encouragement we had the privilege to work.")
    body("We avail this opportunity to thank the respected <Name of HoD>, Head of the Department of "
         "<Name of the Department>, for facilitating such a pleasant environment in the department and also "
         "for providing everlasting encouragement and support throughout.")
    body("We acknowledge with a deep sense of responsibility and gratitude the help rendered by respected "
         "Dr. D. N. Dewangan, Director, RCET, Bhilai; Dr. Chinmay Chandrakar, Executive Director (Academics), "
         "RCET, Bhilai; Dr. Prateek Agrawal, Director, School of CSE and School of IT, RISU, Bhilai; and "
         "Dr. Asheesh Dixit, Director (Academics) and Dean, SOCSE & SOIT, RISU, Bhilai, for infusing endless "
         "enthusiasm and instilling a spirit of dynamism.")
    body("We would also like to thank all faculty members of our department, the supporting staff, and all "
         "the Technical Trainers for always being helpful over the years. Last but not the least, we express "
         "our deepest gratitude to our parents and the management of Rungta College of Engineering and "
         "Technology, Bhilai, for their continuous moral support and encouragement.")
    doc.add_paragraph()
    for i in range(1, 7):
        centered(f"<Student Name {i}, CSVTU Roll No. {i}, Enrollment No. {i}>", 11, space=2)
    page_break()

    # ---- Table of contents ----
    h1("Table of Contents")
    insert_toc()
    page_break()

    # ---- Abstract ----
    h1("Abstract")
    ab = doc.add_paragraph()
    ab.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ab.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    _set_font(ab.add_run(
        "Modern enterprises increasingly depend on containerized, cloud-native infrastructure delivered "
        "through DevSecOps pipelines. While these technologies improve agility, they significantly expand "
        "the attack surface and expose organizations to sophisticated, multi-stage cyber attacks. Many "
        "enterprises, however, lack realistic breach-simulation environments and integrated detection across "
        "distributed container workloads. This project designs, implements, and evaluates an end-to-end "
        "breach lifecycle simulation for a fictional containerized enterprise, “Cryvion Networks”, "
        "built entirely with Docker. A deliberately vulnerable three-tier application is subjected to a "
        "realistic nine-stage attack — reconnaissance, credential brute force, SQL injection, remote code "
        "execution, persistence, privilege-escalation enumeration, credential access, lateral movement, and "
        "data exfiltration — each mapped to the MITRE ATT&CK framework and the Cyber Kill Chain. A Wazuh "
        "SIEM with custom, ATT&CK-tagged detection rules and file-integrity monitoring detects the campaign, "
        "and an automated collector reconstructs a forensic timeline and indicators of compromise. The "
        "environment is then re-architected following Zero-Trust and least-privilege principles — a "
        "ModSecurity web application firewall, network micro-segmentation, non-root and capability-dropped "
        "containers, secret management, and a DevSecOps CI pipeline. Re-running the identical attack against "
        "the hardened architecture blocked all damaging objectives, demonstrating the effectiveness of "
        "layered detection and hardening. The work bridges offensive simulation, SOC detection engineering, "
        "digital forensics, incident response, and security architecture in a single reproducible laboratory."),
        size=12)
    page_break()

    # ---- Lists ----
    h1("List of Figures")
    for f in [
        "Fig 1.1  Cyber Kill Chain", "Fig 4.1  Vulnerable Architecture",
        "Fig 4.2  Attack Chain mapped to MITRE ATT&CK", "Fig 4.3  Detection Pipeline (Wazuh SIEM)",
        "Fig 4.4  Hardened Architecture (Zero-Trust)", "Fig 4.5  DevSecOps CI/CD Pipeline",
        "Fig 5.1  Lab containers running", "Fig 5.2  Vulnerable customer portal",
        "Fig 5.3  Reconnaissance and brute force", "Fig 5.4  Remote code execution as root",
        "Fig 5.5  Lateral movement and exfiltration", "Fig 5.6  Wazuh dashboard overview",
        "Fig 5.7  Custom ATT&CK-tagged alerts", "Fig 5.8  Wazuh MITRE ATT&CK module",
        "Fig 5.9  ATT&CK Navigator campaign layer", "Fig 5.10  Forensic attack timeline",
        "Fig 5.11  Indicators of compromise", "Fig 5.12  Container filesystem diff",
        "Fig 5.13  Attack blocked on hardened stack", "Fig 5.14  WAF blocking an attack (HTTP 403)",
        "Fig 5.15  Empty loot on hardened stack", "Fig 5.16  DevSecOps gate (Hadolint)",
    ]:
        body(f, justify=False)
    page_break()

    h1("List of Tables")
    for t in [
        "Table 2.1  Comparison of existing breach-simulation / detection approaches",
        "Table 4.1  Container inventory and roles",
        "Table 4.2  Planted vulnerabilities mapped to MITRE ATT&CK",
        "Table 4.3  Custom detection rules mapped to MITRE ATT&CK",
        "Table 4.4  Weakness-to-hardening-control matrix",
        "Table 5.1  Detection results per attack technique",
        "Table 5.2  Before/after attack outcomes (vulnerable vs hardened)",
    ]:
        body(t, justify=False)
    page_break()

    h1("List of Abbreviations")
    abbr = [
        ("ATT&CK", "Adversarial Tactics, Techniques and Common Knowledge"),
        ("SIEM", "Security Information and Event Management"),
        ("HIDS", "Host-based Intrusion Detection System"),
        ("FIM", "File Integrity Monitoring"), ("RCE", "Remote Code Execution"),
        ("SQLi", "SQL Injection"), ("WAF", "Web Application Firewall"),
        ("CRS", "Core Rule Set (OWASP)"), ("IOC", "Indicator of Compromise"),
        ("IR", "Incident Response"), ("NIST", "National Institute of Standards and Technology"),
        ("DMZ", "Demilitarized Zone"), ("CI/CD", "Continuous Integration / Continuous Deployment"),
        ("PII", "Personally Identifiable Information"), ("C2", "Command and Control"),
        ("DAST", "Dynamic Application Security Testing"), ("SAST", "Static Application Security Testing"),
        ("IaC", "Infrastructure as Code"), ("SSH", "Secure Shell"),
        ("WSL", "Windows Subsystem for Linux"),
    ]
    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, v in abbr:
        cells = t.add_row().cells
        _set_font(cells[0].paragraphs[0].add_run(k), size=11, bold=True)
        _set_font(cells[1].paragraphs[0].add_run(v), size=11)
    page_break()


# ===================================================================
# CHAPTERS
# ===================================================================
def chapter1():
    h1("Chapter 1   Introduction")
    h2("1.1  Background")
    body("Modern enterprises increasingly rely on containerized applications, cloud-native deployments, "
         "centralized logging, microservice architectures, and Continuous Integration / Continuous "
         "Deployment (CI/CD) pipelines. Containerization — popularized by Docker — packages an "
         "application together with its dependencies into a lightweight, portable unit, enabling rapid and "
         "consistent deployment. While these technologies improve agility and scalability, they "
         "significantly expand the attack surface.")
    body("Organizations today face sophisticated, multi-stage cyber attacks that combine credential-based "
         "brute-force attempts, SQL injection and web exploitation, container compromise and breakout, "
         "lateral movement across internal networks, data exfiltration through covert channels, and "
         "persistence mechanisms embedded within the infrastructure. Traditional reactive security "
         "approaches are insufficient against such coordinated and evolving threats.")
    h2("1.2  Motivation")
    body("Many enterprises lack realistic breach-simulation environments, integrated detection across "
         "distributed systems, proper visibility into container-level activity, correlated log analytics "
         "for multi-vector attacks, structured incident-response workflows, forensic readiness, and "
         "DevSecOps-based remediation strategies. As a result, they struggle to detect attacks early, "
         "identify root causes, prevent lateral movement, contain compromised containers, and protect "
         "sensitive data. This project addresses that gap by building a complete, reproducible laboratory "
         "in which a breach can be safely executed, detected, investigated, and ultimately prevented.")
    h2("1.3  Objectives")
    for o in [
        "Design and build a realistic containerized enterprise (web application, database, internal service, and SIEM).",
        "Simulate a multi-stage breach spanning the full Cyber Kill Chain.",
        "Detect each stage using a SIEM, custom rules, and runtime/host monitoring.",
        "Investigate the breach forensically — recover artifacts, reconstruct a timeline, and extract IOCs.",
        "Map the entire attack chain to the MITRE ATT&CK framework.",
        "Define a structured incident-response lifecycle aligned with NIST SP 800-61.",
        "Re-architect the environment using Zero-Trust hardening and a DevSecOps CI pipeline.",
        "Communicate findings through an executive-level incident report.",
    ]:
        numbered(o)
    h2("1.4  Scope and Limitations")
    body("The work is conducted in a controlled laboratory built on Docker Desktop (WSL2) on a single host. "
         "All offensive activity is confined to private Docker networks and never targets external systems. "
         "The environment models an enterprise at small scale; it is not a production deployment. A genuine "
         "container-to-host breakout via the Docker socket is identified and demonstrated as a vector but "
         "treated as a stretch objective. Network-layer intrusion detection (NIDS) is identified as future "
         "work.")
    figure(os.path.join(FIG, "fig1_1_killchain.png"), "Fig 1.1  Cyber Kill Chain")
    h2("1.5  Organization of the Report")
    body("Chapter 2 reviews the relevant frameworks and technologies. Chapter 3 identifies the problem and "
         "derives requirements. Chapter 4 details the methodology — architecture, attack simulation, "
         "detection, forensics, incident response, hardening, and CI. Chapter 5 presents the results, and "
         "Chapter 6 concludes with limitations and future work.")
    page_break()


def chapter2():
    h1("Chapter 2   Literature Review")
    h2("2.1  The Cyber Kill Chain")
    body("The Cyber Kill Chain, introduced by Lockheed Martin, models an intrusion as a sequence of stages "
         "— reconnaissance, weaponization, delivery, exploitation, installation, command and control, "
         "and actions on objectives. Understanding these stages allows defenders to break the chain at the "
         "earliest possible point. This project structures its simulated attack along an equivalent "
         "progression from external reconnaissance to data exfiltration.")
    h2("2.2  The MITRE ATT&CK Framework")
    body("MITRE ATT&CK is a globally accessible knowledge base of adversary tactics and techniques based on "
         "real-world observations. Tactics describe the attacker’s goal (e.g., Initial Access, Lateral "
         "Movement), while techniques (e.g., T1190 Exploit Public-Facing Application) describe how the goal "
         "is achieved. ATT&CK provides a common language for mapping attacks and validating detection "
         "coverage, and is used throughout this project to label both the attack and the detections.")
    h2("2.3  NIST SP 800-61 Incident-Response Lifecycle")
    body("NIST Special Publication 800-61 defines a four-phase incident-handling lifecycle — "
         "Preparation; Detection and Analysis; Containment, Eradication and Recovery; and Post-Incident "
         "Activity. This project applies the lifecycle to the simulated breach, demonstrating "
         "container-native containment and recovery onto a hardened architecture.")
    h2("2.4  Container Security Fundamentals")
    body("Containers share the host kernel and are isolated using Linux namespaces and control groups "
         "(cgroups), with privileges governed by capabilities. Common misconfigurations — running as "
         "root, retaining all capabilities, mounting the Docker socket, and using writable root filesystems "
         "— weaken this isolation and can enable privilege escalation or container escape. These "
         "anti-patterns are deliberately introduced in the vulnerable design and removed in the hardened one.")
    h2("2.5  SIEM and Host-Based Detection")
    body("A Security Information and Event Management (SIEM) platform centralizes and correlates logs to "
         "detect threats. Wazuh is an open-source SIEM combining host-based intrusion detection, log "
         "analysis, and file-integrity monitoring; it comprises agents, a manager (with the analysisd "
         "detection engine and custom rules/decoders), an OpenSearch-based indexer, and a dashboard. Wazuh "
         "rules can be tagged with MITRE ATT&CK technique identifiers. Runtime tools such as Falco provide "
         "syscall-level container visibility.")
    h2("2.6  Web Application Security and the WAF")
    body("The OWASP Top 10 enumerates the most critical web-application risks, with injection (including "
         "SQL injection) and security misconfiguration among the most damaging. A Web Application Firewall "
         "(WAF) such as ModSecurity, configured with the OWASP Core Rule Set (CRS), inspects HTTP traffic "
         "and blocks malicious requests using an anomaly-scoring model governed by configurable paranoia "
         "levels. In this project the WAF serves as a compensating control alongside fixed application code.")
    h2("2.7  Zero Trust and Micro-Segmentation")
    body("Zero Trust assumes no implicit trust between network segments and enforces least-privilege "
         "communication. Micro-segmentation divides the environment into small zones so that a compromise of "
         "one component cannot freely reach others. The hardened architecture applies this principle so that "
         "the web tier can reach only the database it requires and never the internal host.")
    h2("2.8  DevSecOps Tooling")
    body("DevSecOps integrates security into the CI/CD pipeline through automated gates: static application "
         "security testing (SAST, e.g., Semgrep), Dockerfile linting (Hadolint), infrastructure-as-code "
         "scanning (Checkov), container image vulnerability scanning (Trivy), and dynamic application "
         "security testing (DAST, e.g., OWASP ZAP). Each gate can fail the build, preventing insecure "
         "artifacts from shipping.")
    h2("2.9  Related Work and Research Gap")
    body("Existing resources address parts of the problem in isolation. Damn Vulnerable Web Application "
         "(DVWA) provides a vulnerable target but no detection or hardening. Atomic Red Team and MITRE "
         "Caldera automate adversary techniques but are not packaged as a containerized enterprise with an "
         "integrated SIEM and a hardened counterpart. The gap this project fills is an end-to-end, "
         "reproducible lifecycle — attack, detection, forensics, response, and hardening — within a "
         "single containerized environment.")
    table("Table 2.1  Comparison of existing breach-simulation / detection approaches",
          ["Approach", "Attack", "Detection", "Hardening", "Containerized"],
          [["DVWA", "Yes", "No", "No", "Partial"],
           ["Atomic Red Team", "Yes", "Partial", "No", "No"],
           ["MITRE Caldera", "Yes", "Partial", "No", "No"],
           ["This project (Cryvion)", "Yes", "Yes (Wazuh+MITRE)", "Yes (Zero-Trust+CI)", "Yes"]])
    page_break()


def chapter3():
    h1("Chapter 3   Problem Identification")
    h2("3.1  The Core Problem")
    body("Enterprises operating containerized, DevSecOps-driven infrastructure are exposed to coordinated, "
         "multi-stage attacks, yet frequently lack the means to study such attacks safely or to validate "
         "their detection and response capabilities. Without a realistic environment, detection rules go "
         "untested, forensic readiness is unknown, and architectural weaknesses remain hidden until a real "
         "incident occurs.")
    h2("3.2  Specific Gaps")
    for g in [
        "Absence of realistic, repeatable breach-simulation environments for containerized systems.",
        "Limited visibility into container-level activity (process execution, file changes, in-container logins).",
        "Weak correlation of logs across distributed components for multi-vector attacks.",
        "Inadequate controls against lateral movement in flat internal networks.",
        "Poor forensic readiness in ephemeral container environments.",
        "Lack of DevSecOps gates to prevent insecure configurations from being deployed.",
    ]:
        bullet(g)
    h2("3.3  Problem Statement")
    body("Design and evaluate an end-to-end breach lifecycle simulation in a containerized environment that "
         "(a) reproduces a realistic multi-stage attack, (b) detects and investigates it using a SIEM and "
         "forensic tooling, (c) maps it to industry frameworks, and (d) demonstrates a hardened "
         "re-architecture that prevents recurrence — all within a safe, reproducible laboratory.")
    h2("3.4  Derived Requirements")
    for r in [
        "A multi-tier containerized target with realistic, intentional vulnerabilities.",
        "An automated, repeatable attack covering the full kill chain.",
        "Centralized detection with ATT&CK-tagged, high-fidelity alerts.",
        "Automated forensic collection producing a timeline and IOCs.",
        "A hardened architecture and DevSecOps pipeline, validated by re-running the attack.",
    ]:
        numbered(r)
    page_break()


def chapter4():
    h1("Chapter 4   Methodology")
    h2("4.1  System Architecture")
    body("The environment models a small enterprise, “Cryvion Networks”, composed entirely of Docker "
         "containers and divided into network zones that emulate real enterprise segmentation. A web "
         "application is reachable from a demilitarized zone (DMZ); a database and an internal service "
         "reside on an internal network; and a Wazuh SIEM monitors all components. The web application is "
         "intentionally dual-homed across the DMZ and internal networks, so that a single web compromise can "
         "bridge into the internal trust zone — the central design flaw that the attack exploits.")
    figure(os.path.join(FIG, "fig4_1_vuln_arch.png"), "Fig 4.1  Vulnerable Architecture (the “before” state)")
    h3("4.1.1  Component Inventory")
    table("Table 4.1  Container inventory and roles",
          ["Container", "Technology", "Network(s)", "Role"],
          [["web-app", "Custom Flask app", "dmz, internal", "Public customer portal; entry point"],
           ["db", "MySQL 8", "internal", "Customer PII (crown jewels)"],
           ["internal-svc", "OpenSSH host", "internal", "Internal backup host; lateral target"],
           ["attacker", "Debian + tooling", "dmz", "External adversary"],
           ["Wazuh stack", "Manager/Indexer/Dashboard", "monitoring", "SIEM and detection"]])
    h2("4.2  Vulnerable Environment Design")
    h3("4.2.1  Web Application and Planted Vulnerabilities")
    body("The customer portal is a custom Python/Flask application deliberately containing several classic "
         "vulnerabilities, each chosen to exercise a specific ATT&CK technique. It runs as root and includes "
         "post-exploitation tooling, modelling a poorly hardened image.")
    table("Table 4.2  Planted vulnerabilities mapped to MITRE ATT&CK",
          ["Vulnerability", "Location", "MITRE"],
          [["SQL injection (auth bypass)", "/login", "T1190"],
           ["SQL injection (UNION data theft)", "/dashboard", "T1190 / T1005"],
           ["OS command injection (RCE)", "/admin/diagnostics", "T1059"],
           ["Unrestricted file upload", "/upload", "T1505"],
           ["Plaintext / reused credentials", "users table, db.conf", "T1552"],
           ["Container runs as root + Docker socket mounted", "Dockerfile / compose", "T1611"]])
    code_block(
        "# /login — string-concatenated SQL (injectable)\n"
        "q = \"SELECT * FROM users WHERE username='%s' AND password='%s'\" % (u, p)\n"
        "# /admin/diagnostics — shell=True with user input (RCE)\n"
        "subprocess.run('ping -c 1 ' + host, shell=True, capture_output=True)")
    h3("4.2.2  Database and Internal Host")
    body("The MySQL database is seeded with synthetic customer PII (names, emails, SSNs, card numbers) and "
         "weak user passwords that drive the brute-force stage. The internal service is an SSH host holding "
         "a backup export and a flag; it shares a service-account credential with the web tier, enabling "
         "credential-reuse-based lateral movement.")
    h2("4.3  Attack Simulation")
    body("The attacker begins with reach only into the DMZ. Every action deeper in the network is therefore "
         "driven through the compromised web tier, validating the segmentation and pivot model. The full "
         "nine-stage chain is orchestrated by a single script and mapped to MITRE ATT&CK as shown below.")
    figure(os.path.join(FIG, "fig4_2_attack_chain.png"), "Fig 4.2  Attack Chain mapped to MITRE ATT&CK")
    stages = [
        ("4.3.1  Reconnaissance (T1595)", "An nmap service scan identifies the exposed web port and fingerprints the application."),
        ("4.3.2  Initial Access (T1110 / T1190)", "Hydra brute-forces the login form and recovers admin/admin123; a SQL-injection payload (admin' -- -) bypasses authentication to establish a session."),
        ("4.3.3  Execution (T1059 / T1505)", "Command injection in the diagnostics endpoint yields remote code execution as root; an unrestricted upload drops a web shell."),
        ("4.3.4  Persistence (T1136 / T1053 / T1098)", "A uid-0 backdoor account, a malicious cron beacon, and an attacker SSH key are planted."),
        ("4.3.5  Privilege-Escalation Enumeration (T1611)", "The mounted Docker socket is identified as a container-escape vector to the host."),
        ("4.3.6  Credential Access (T1552)", "Database and internal-host credentials are harvested from environment variables and an on-disk configuration file."),
        ("4.3.7  Lateral Movement (T1021 / T1078)", "Using the reused credential, the attacker pivots over SSH from the web container to the internal backup host."),
        ("4.3.8  Collection (T1005)", "A UNION-based SQL injection dumps user credentials and PII; the internal backup export and flag are stolen via the pivot."),
        ("4.3.9  Exfiltration (T1041 / T1048)", "Stolen data is sent to an attacker-controlled command-and-control listener over HTTP, with a DNS-tunnelling demonstration."),
    ]
    for title, desc in stages:
        h3(title); body(desc)
    h2("4.4  Detection and Monitoring")
    body("Wazuh agents are deployed inside the web and internal containers; they perform file-integrity "
         "monitoring and forward application and authentication logs to the manager. Because containers log "
         "to files rather than the systemd journal, the agents are configured to tail the application "
         "security log and the SSH authentication log explicitly. The application emits structured security "
         "events that a custom decoder parses and custom rules classify and tag with MITRE techniques.")
    figure(os.path.join(FIG, "fig4_3_detection.png"), "Fig 4.3  Detection Pipeline (Wazuh SIEM)")
    table("Table 4.3  Custom detection rules mapped to MITRE ATT&CK",
          ["Rule ID", "Detects", "MITRE"],
          [["100110", "SQL injection", "T1190"],
           ["100111", "Command injection / RCE", "T1059"],
           ["100112 / 100113", "Failed login / brute-force correlation", "T1110"],
           ["100114", "File upload (web shell)", "T1505"],
           ["550 (FIM)", "File-integrity change (persistence)", "T1053"],
           ["5715 (sshd)", "Lateral SSH authentication", "T1021 / T1078"]])
    h2("4.5  Forensic Methodology")
    body("An automated collector captures, for each compromised container, the filesystem diff against the "
         "immutable base image, container logs, persistence artifacts, process and network state, and the "
         "application and authentication logs. It also extracts the attacker loot, reconstructs an attack "
         "timeline from SIEM alerts, derives indicators of compromise, and bundles the evidence. Because the "
         "container image is immutable, the filesystem diff isolates exactly what the attacker introduced — "
         "a key advantage of container forensics.")
    h2("4.6  Incident-Response Lifecycle")
    body("The breach is handled following NIST SP 800-61. Preparation comprises the deployed SIEM, defined "
         "segmentation, and a ready collector. Detection and Analysis rely on the ATT&CK-tagged alerts and "
         "the reconstructed timeline. Containment uses container-native actions — disconnecting the "
         "compromised container from its networks and pausing it for live forensics. Eradication removes "
         "persistence and rebuilds from the clean image, and Recovery redeploys onto the hardened "
         "architecture. Post-incident activity feeds the lessons learned into the hardened design and CI "
         "pipeline.")
    h2("4.7  Hardened Re-Architecture")
    body("Every root cause is addressed by at least one control, and lateral movement is contained by "
         "network micro-segmentation. The web tier is fronted by a ModSecurity WAF and can reach only the "
         "database; the internal host is isolated on its own segment. Containers run as non-root with all "
         "capabilities dropped, no new privileges, and read-only filesystems; the Docker socket is removed; "
         "and the database password is provided through a Docker secret. The application code itself is fixed "
         "to use parameterized queries and non-shell command execution.")
    figure(os.path.join(FIG, "fig4_4_hardened.png"), "Fig 4.4  Hardened Architecture (Zero-Trust, the “after” state)")
    table("Table 4.4  Weakness-to-hardening-control matrix",
          ["Weakness", "Hardening control", "Blocks stage"],
          [["SQLi / command injection", "Parameterized queries, no-shell exec, WAF", "Initial access, RCE"],
           ["Flat, dual-homed network", "Zero-Trust micro-segmentation", "Lateral movement"],
           ["Container runs as root", "Non-root user", "Privilege escalation"],
           ["All capabilities / writable FS", "cap_drop ALL, read-only FS, no-new-privileges", "Privilege escalation"],
           ["Docker socket mounted", "Socket removed", "Container breakout"],
           ["Plaintext / reused credentials", "Docker secrets, key-only SSH", "Credential access"],
           ["No release security gates", "DevSecOps CI pipeline", "Insecure deployment"]])
    h2("4.8  DevSecOps CI/CD Pipeline")
    body("A continuous-integration pipeline introduces shift-left security gates so that a vulnerable image "
         "or configuration cannot ship. Dockerfile linting, static analysis, infrastructure-as-code "
         "scanning, image vulnerability scanning, and dynamic application testing each fail the build on "
         "high-severity findings.")
    figure(os.path.join(FIG, "fig4_5_devsecops.png"), "Fig 4.5  DevSecOps CI/CD Pipeline")
    page_break()


def chapter5():
    h1("Chapter 5   Results")
    h2("5.1  Breach Execution")
    body("The complete environment was built and verified to run, and the nine-stage attack executed "
         "successfully against the vulnerable stack. Reconnaissance fingerprinted the portal; Hydra "
         "recovered admin/admin123; command injection produced a root shell; persistence was planted; the "
         "attacker pivoted to the internal host and captured the flag; and customer PII together with the "
         "internal backup was exfiltrated to the attacker’s command-and-control listener.")
    screenshot("Fig 5.1  Lab containers running", "Hint: output of `docker compose ps` (all containers Up).")
    screenshot("Fig 5.2  Vulnerable customer portal", "Hint: browser at http://localhost:8080.")
    screenshot("Fig 5.3  Reconnaissance and brute force", "Hint: run-attack stages 1–2 (nmap + Hydra cracking admin/admin123).")
    screenshot("Fig 5.4  Remote code execution as root", "Hint: stage 3 terminal showing uid=0(root).")
    screenshot("Fig 5.5  Lateral movement and exfiltration", "Hint: stages 7–9 (PIVOT_OK, FLAG captured, C2 receives data).")
    h2("5.2  Detection Results")
    body("The SIEM detected the campaign across the kill chain. Custom ATT&CK-tagged rules fired on SQL "
         "injection, command-execution, brute force (including a correlation alert), and web-shell upload; "
         "file-integrity monitoring flagged the persistence changes; and the built-in sshd rule detected the "
         "lateral-movement login. Six distinct ATT&CK techniques were observed.")
    table("Table 5.1  Detection results per attack technique",
          ["Attack action", "Rule", "MITRE", "Detected"],
          [["Brute force", "100112 / 100113", "T1110", "Yes"],
           ["SQL injection", "100110", "T1190", "Yes"],
           ["Command injection / RCE", "100111", "T1059", "Yes"],
           ["Web-shell upload", "100114 / FIM 550", "T1505", "Yes"],
           ["Persistence (file change)", "FIM 550", "T1053", "Yes"],
           ["Lateral movement (SSH)", "5715", "T1021 / T1078", "Yes"],
           ["Reconnaissance / network exfil", "—", "T1595 / T1048", "Gap (no NIDS)"]])
    screenshot("Fig 5.6  Wazuh dashboard overview", "Hint: https://localhost:443 security events overview.")
    screenshot("Fig 5.7  Custom ATT&CK-tagged alerts", "Hint: filter rule IDs 100110–100115 in the events view.")
    screenshot("Fig 5.8  Wazuh MITRE ATT&CK module", "Hint: Wazuh → MITRE ATT&CK module.")
    screenshot("Fig 5.9  ATT&CK Navigator campaign layer", "Hint: load mitre/attack-navigator-layer.json in the ATT&CK Navigator.")
    h2("5.3  Forensic Findings")
    body("The forensic collector reconstructed a chronological, ATT&CK-tagged attack timeline and extracted "
         "indicators of compromise — the web-shell SHA-256 hash, the uid-0 backdoor account, the cron "
         "beacon, the rogue SSH key, the attacker source IP, and the command-and-control endpoint. The "
         "container filesystem diff isolated precisely the files the attacker introduced.")
    screenshot("Fig 5.10  Forensic attack timeline", "Hint: forensics/output/<ts>/attack-timeline.txt.")
    screenshot("Fig 5.11  Indicators of compromise", "Hint: forensics/output/<ts>/IOCs.md.")
    screenshot("Fig 5.12  Container filesystem diff", "Hint: cryvion-web/docker-diff.txt (added/changed files).")
    h2("5.4  Hardening Validation")
    body("The identical attack was re-run against the hardened architecture. All damaging objectives were "
         "blocked: the SQL-injection bypass returned “Invalid credentials”, remote code execution and "
         "data theft were prevented (producing empty loot files), lateral movement failed due to "
         "segmentation, the Docker socket was absent, and the on-disk credentials file no longer existed — "
         "while legitimate traffic continued to function.")
    table("Table 5.2  Before/after attack outcomes (vulnerable vs hardened)",
          ["Attacker objective", "Vulnerable", "Hardened"],
          [["SQLi auth bypass", "Succeeds", "Blocked (Invalid credentials)"],
           ["SQLi data theft", "Dumps PII", "Blocked (param query + WAF 403)"],
           ["Command injection / RCE", "Root shell", "Blocked (empty loot + WAF 403)"],
           ["PII exfiltration", "5 SSNs stolen", "Blocked (0 bytes)"],
           ["Credential harvest", "db.conf dumped", "Blocked (file removed)"],
           ["Lateral movement", "PIVOT_OK", "Blocked (segmentation)"],
           ["Container breakout", "Socket present", "No socket found"]])
    screenshot("Fig 5.13  Attack blocked on hardened stack", "Hint: hardened run-attack terminal (failures).")
    screenshot("Fig 5.14  WAF blocking an attack (HTTP 403)", "Hint: curl SQLi/RCE returning 403.")
    screenshot("Fig 5.15  Empty loot on hardened stack", "Hint: `ls -l attacker/loot-hardened/` (0-byte files).")
    h2("5.5  DevSecOps Results")
    body("The CI gates were demonstrated locally: Hadolint flagged the vulnerable Dockerfile (unpinned "
         "packages, root user) and passed the hardened one, confirming that the pipeline would have "
         "prevented the insecure image from shipping.")
    screenshot("Fig 5.16  DevSecOps gate (Hadolint)", "Hint: Hadolint output on vulnerable vs hardened Dockerfile.")
    h2("5.6  Discussion")
    body("The results show that a poorly configured containerized enterprise can be fully compromised "
         "through a chained set of common weaknesses, but that layered detection makes the attack visible at "
         "every stage and that Zero-Trust hardening combined with code fixes contains it completely. The "
         "principal limitation is the absence of network-layer detection, which left reconnaissance and raw "
         "network exfiltration unobserved at the network tier; this is addressed in future work.")
    page_break()


def chapter6():
    h1("Chapter 6   Conclusion")
    h2("6.1  Summary")
    body("This project delivered a complete, reproducible breach lifecycle for a containerized enterprise: a "
         "vulnerable multi-tier environment, an automated nine-stage attack mapped to MITRE ATT&CK, a Wazuh "
         "SIEM with custom ATT&CK-tagged detection, automated forensics, a NIST-aligned incident-response "
         "model, and a hardened Zero-Trust re-architecture with a DevSecOps pipeline. Running the same "
         "attack against both architectures produced opposite outcomes — a full breach versus complete "
         "containment — demonstrating the value of detection engineering and least-privilege hardening.")
    h2("6.2  Limitations")
    for l in [
        "No network intrusion detection system, so reconnaissance and raw network exfiltration are not detected at the network layer.",
        "The environment runs on a single host and at small scale; it does not model orchestration (e.g., Kubernetes).",
        "Application weak-password policy is mitigated by the WAF and code fixes but not by an account-lockout / MFA mechanism.",
    ]:
        bullet(l)
    h2("6.3  Future Work")
    for f in [
        "Deploy a network IDS (Suricata) to close the reconnaissance and exfiltration visibility gap.",
        "Add runtime syscall monitoring (Falco) for deeper container-level detection.",
        "Extend the lab to Kubernetes to study orchestration-level attacks and policies.",
        "Enforce authentication policy (account lockout, multi-factor authentication).",
        "Automate the full DevSecOps pipeline in a hosted CI service with policy-as-code gates.",
    ]:
        bullet(f)
    page_break()


def back_matter():
    h1("Bibliography")
    refs = [
        "MITRE Corporation, “MITRE ATT&CK® Enterprise Matrix”, attack.mitre.org, 2024.",
        "Lockheed Martin, “The Cyber Kill Chain®”, lockheedmartin.com, 2011.",
        "P. Cichonski et al., “Computer Security Incident Handling Guide”, NIST SP 800-61 Rev. 2, 2012.",
        "Wazuh Inc., “Wazuh – The Open Source Security Platform: Documentation”, documentation.wazuh.com, 2024.",
        "OWASP Foundation, “OWASP Top 10:2021”, owasp.org/Top10, 2021.",
        "OWASP Foundation, “OWASP ModSecurity Core Rule Set”, coreruleset.org, 2024.",
        "Docker Inc., “Docker Documentation – Security”, docs.docker.com/engine/security, 2024.",
        "Aqua Security, “Trivy: Vulnerability and Misconfiguration Scanner”, trivy.dev, 2024.",
        "NIST, “Zero Trust Architecture”, NIST SP 800-207, 2020.",
        "Falco Authors, “Falco – Cloud Native Runtime Security”, falco.org, 2024.",
    ]
    for i, r in enumerate(refs, 1):
        p = doc.add_paragraph(f"[{i}]  {r}")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    page_break()

    h1("Appendix")
    h2("A.  Custom Wazuh Detection Rule (excerpt)")
    code_block(
        "<rule id=\"100110\" level=\"10\">\n"
        "  <if_sid>100100</if_sid>\n"
        "  <field name=\"crvevent\">SQLI</field>\n"
        "  <description>SQL injection attempt on CryvionPortal: $(detail)</description>\n"
        "  <mitre><id>T1190</id></mitre>\n"
        "</rule>")
    h2("B.  Hardened Container Configuration (excerpt)")
    code_block(
        "app:\n"
        "  read_only: true\n"
        "  cap_drop: [ALL]\n"
        "  security_opt: [\"no-new-privileges:true\"]\n"
        "  secrets: [db_password]\n"
        "  networks: [app_net, data_net]   # not on edge/mgmt")
    h2("C.  Attack Orchestrator (excerpt)")
    code_block(
        "bash 01-recon.sh; bash 02-initial-access.sh; bash 03-execution.sh\n"
        "bash 04-persistence.sh; bash 05-privesc.sh; bash 06-credaccess.sh\n"
        "bash 07-lateral.sh; bash 08-collection.sh; bash 09-exfil.sh")
    body("The complete source code, configuration, detection rules, and the ATT&CK Navigator layer are "
         "maintained in the project repository.")


# ===================================================================
# BUILD
# ===================================================================
front_matter()
_pgnum(doc.sections[0], "lowerRoman", start=1)
add_page_number_footer(doc.sections[0])

# new section for the body with decimal page numbers restarting at 1
doc.add_section(WD_SECTION.NEW_PAGE)
body_sec = doc.sections[-1]
body_sec.page_width, body_sec.page_height = Mm(210), Mm(297)
body_sec.left_margin = Inches(1.3)
body_sec.right_margin = body_sec.top_margin = body_sec.bottom_margin = Inches(1.0)
_pgnum(body_sec, "decimal", start=1)
add_page_number_footer(body_sec)

chapter1(); chapter2(); chapter3(); chapter4(); chapter5(); chapter6(); back_matter()

doc.save(OUT)
print("REPORT WRITTEN ->", OUT)
print("paragraphs:", len(doc.paragraphs), " tables:", len(doc.tables))
